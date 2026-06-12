# -*- coding: utf-8 -*-
"""生成两份提交材料（“三智能体”AI赋能版）：
1. 师范生教育教学能力认定抽检评价改革实践-案例文稿.docx
2. 附件1 教育评价领域典型实践案例征集表.docx
图1~图4 以占位行标出，由用户自行插入图片。
"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))

CASE_TITLE_MAIN = "三智能体协同　智赋评价"
CASE_TITLE_SUB = "——师范生教育教学能力认定抽检的AI赋能全流程质量监测实践"
CASE_NAME_ONELINE = "三智能体协同 智赋评价——师范生教育教学能力认定抽检的AI赋能全流程质量监测实践"


def set_font(run, cn_font="仿宋", en_font="Times New Roman", size=14, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = en_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), cn_font)
    rfonts.set(qn('w:ascii'), en_font)
    rfonts.set(qn('w:hAnsi'), en_font)
    if color is not None:
        run.font.color.rgb = color


# ============================================================
# 文档一：案例文稿
# ============================================================

def build_case_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.6)

    def title():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(CASE_TITLE_MAIN)
        set_font(r, cn_font="方正小标宋简体", size=20, bold=True)
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(16)
        r2 = p2.add_run(CASE_TITLE_SUB)
        set_font(r2, cn_font="黑体", size=15, bold=False)

    def h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.5
        r = p.add_run(text)
        set_font(r, cn_font="黑体", size=15, bold=True)

    def h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.5
        r = p.add_run(text)
        set_font(r, cn_font="黑体", size=14, bold=True)

    def h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(28)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.line_spacing = 1.5
        r = p.add_run(text)
        set_font(r, cn_font="仿宋", size=14, bold=True)

    def body(text):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(28)
        p.paragraph_format.line_spacing = 1.5
        r = p.add_run(text)
        set_font(r, cn_font="仿宋", size=14)

    def figure(caption):
        ph = doc.add_paragraph()
        ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ph.paragraph_format.space_before = Pt(6)
        ph.paragraph_format.space_after = Pt(2)
        rr = ph.add_run("［请在此处插入图片］")
        set_font(rr, cn_font="仿宋", size=12, color=RGBColor(0x99, 0x99, 0x99))
        pc = doc.add_paragraph()
        pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pc.paragraph_format.space_after = Pt(8)
        rc = pc.add_run(caption)
        set_font(rc, cn_font="楷体", size=12, bold=True)

    title()

    # ---------- 一、案例背景 ----------
    h1("一、案例背景")
    body("师范生教育教学能力认定是教师教育质量保障的重要环节，抽检工作直接关系师范生培养质量监测的科学性、公平性和公信力。"
         "按照《重庆市2025年师范生教育教学能力认定抽检实施方案（试行）》要求，抽检工作需覆盖样本抽取、材料上传、校级审核、"
         "重庆市教育评估院审核、专家评分、结果发布等多个环节，涉及重庆市教育评估院、高校、学生、专家等多类主体，"
         "业务链条长、协同难度大、规范要求高。")
    body("以往类似工作主要依赖线下通知、人工汇总和经验判断，存在三方面突出问题。"
         "一是标准执行不统一，政策文件中的指标体系、评分规则、举证要求较为复杂，缺乏数字化承载，各校理解和执行容易产生偏差。"
         "二是审核评审效率低，传统系统对材料的校验仅停留在数量、格式、大小等规则层面，内容是否符合要求仍依赖人工逐份打开查看；"
         "专家需阅读大量PDF和视频材料并撰写结构化意见，单个学生评阅耗时较长，且容易出现意见趋同化、模板化等问题。"
         "三是结果运用支撑弱，评审结果多以分数明细形式呈现，缺乏结构化诊断分析，学校难以从中获取有价值的改进信息，"
         "评价的诊断功能和反哺功能未能充分释放。")
    body("基于上述问题，项目团队提出“把政策要求转化为平台规则，把人工核查升级为智能预检，把专家评审升级为人机协同，"
         "把结果汇总升级为智能诊断”的总体思路，研发建设师范生教育教学能力认定抽检系统，并创新嵌入三大AI智能体，"
         "以数字化与智能化双轮驱动，推动抽检工作从经验驱动走向规则驱动、从人工核查走向人机协同、从结果管理走向智能诊断。")
    figure("图1　系统总体架构与三智能体业务流程图")

    # ---------- 二、具体举措及实施成效 ----------
    h1("二、具体举措及实施成效")
    h2("（一）构建“三智能体”AI赋能体系，破解评价全链条共性难题")
    body("本案例最具突破性的创新在于围绕抽检业务全链条，精准嵌入三个差异化AI智能体，分别赋能学生与学校审核、专家评审、"
         "结果反馈三大关键环节，形成“贯通上下游、人机各司其职”的智能化质量监测体系。")

    h3("1.佐证材料“智能预检”，赋能学生上传与学校管理员审核")
    body("针对传统抽检系统仅做规则层面校验、内容判断仍依赖人工的突出问题，项目团队在学生上传与校级审核环节嵌入佐证材料智能预检智能体，"
         "构建“双场景预检”机制。在学生提交前的自助预检场景下，学生上传举证材料后可一键启动智能预检，智能体基于该指标对应的评价标准和举证要求，"
         "自动检测材料的内容相关性、要素完整性、表述规范性、内容深度以及多文档一致性等多个维度，并生成可视化预检报告，"
         "学生可据此反复修改、迭代优化，待材料质量达标后再正式提交。在校管审核时的辅助审阅场景下，校级管理员进入审核界面时系统自动呈现该学生材料的智能预检结果，"
         "重点问题以红、黄、绿三色风险标识直观呈现，支持快速定位到具体问题位置（如某PDF的第几页、某段落），校管可在AI辅助下高效完成审阅判断，"
         "最终的“通过、驳回”决定权仍完全保留在校管手中，AI仅作为“审阅助手”提供线索和参考，避免技术替代人的专业判断。")
    figure("图2　佐证材料“智能预检”反馈界面截图")

    h3("2.专家评审“AI辅助”，赋能专家评阅与意见撰写")
    body("针对专家评审环节材料阅读量大、意见撰写耗时、质量保障难等突出问题，项目团队在专家评审环节嵌入评审辅助智能体，"
         "提供“评审参考意见智能生成”服务。当专家打开某学生的评审页面时，评审辅助智能体基于该学生的完整举证材料、工具库下的认定标准与评分规则、"
         "各二级指标的考察要点综合分析，自动生成针对该学生、针对每条二级指标的结构化初步评审意见，包含指标达成度初判、关键证据定位、优势点提炼、"
         "不足点提示、意见建议参考等要素。在人机协同设计上，AI生成的意见明确标注为“参考建议”，专家可采纳、修改、补充或完全自主撰写，"
         "评分和最终意见均由专家独立填写并确认，AI不参与最终评分判定，系统记录AI生成内容与专家最终内容的差异，确保评审过程可追溯。")
    figure("图3　专家评审界面（含AI辅助意见生成）")

    h3("3.评审结果“智能诊断报告”，赋能学校结果反馈与管理决策")
    body("针对抽检结果传统呈现方式难以支撑学校精准改进的突出问题，项目团队构建认定结果分析智能体，在某一工具库评审基本完成后自动激活，"
         "基于全量评审数据生成结构化诊断报告。报告内容涵盖总体表现概览、能力维度画像、指标达成分析、典型问题归纳、横向对比定位、改进建议参考等核心板块，"
         "形成“诊断书加处方”式的综合报告。对学校而言，收到的不再是单一分数，而是能够直接对接师范生培养方案优化、实践教学改进、薄弱指标专项强化等具体工作的诊断建议；"
         "对重庆市教育评估院而言，基于全市抽检数据，智能体可自动汇总各校诊断报告共性发现，为重庆市教育评估院形成全市师范生培养质量年度分析提供高质量数据支撑。")
    figure("图4　评审结果学校维度、管理端“智能诊断报告”展示界面截图")

    h2("（二）推动标准数字化，解决“评价要求落不细”问题")
    body("为夯实AI智能化运行的基础底座，项目团队系统推动评价标准的数字化转化。一是将抽检标准转化为可配置指标表，"
         "系统围绕抽检标准建立指标表管理机制，支持分层配置评价指标及评分规则，将原本停留在文本中的认定标准转化为平台中的可执行规则，"
         "这一基础也为三个AI智能体提供了核心规则依据。二是将评分办法转化为可计算规则，系统支持等级制和分数制两种评价方式，"
         "等级制可设置选项对应分值，实现定性判断与定量换算相结合，既保留教育评价的专业判断空间，又为认定结果智能诊断提供了可计算、可分析的数据基础。"
         "三是将举证要求转化为材料清单，系统建立举证材料库，支持PDF、MP4等格式上传，并对数量、大小、上传主体、排序等属性进行精细化设置，"
         "通过指标与举证材料多对多关联，把“用什么材料证明什么能力”清晰呈现出来，也为佐证材料智能预检提供了精准的检测靶向。")

    h2("（三）推动流程闭环化，解决“多方协同效率低”问题")
    body("为破解多方协同效率低的难题，项目团队系统重构业务流程。一是构建多角色分权协同机制，系统设置重庆市教育评估院管理员、学校管理员、学生、专家等角色，"
         "按职责赋予不同权限，学生负责上传材料并接受AI预检辅助，学校管理员负责初审和补充指定材料并使用AI辅助审阅，"
         "重庆市教育评估院管理员负责复审、抽检组织和结果管理，专家负责在线评分与评审意见填写并使用AI辅助生成意见，"
         "各角色边界清晰、责任明确，AI在每个角色对应环节提供差异化辅助。二是建立分级审核链条，抽中的学生上传材料后先由校级管理员审核，"
         "再进入重庆市教育评估院管理员审核，审核通过后流转至专家评分环节，将“谁提交、谁审核、谁把关”清晰嵌入系统流程。"
         "三是实现在线查证与全过程留痕，专家可直接在线查看与指标关联的佐证材料，结合评分界面完成评价并填写意见，"
         "系统同步记录登录日志、操作日志、AI辅助使用记录和流程节点信息，确保关键操作可查询、可追溯、可复盘。")

    h2("（四）推动规则公平化，解决“抽样与评审难统一”问题")
    body("为保障抽检过程的公平性和评审组织的灵活性，项目团队创新规则设计。一是创新样本抽取规则，系统支持按专业人数动态设置抽取比例，"
         "并按同校同专业学生“按序等距抽取”，兼顾随机性、代表性和操作公平，减少了人工抽样带来的随意性；"
         "抽取后重庆市教育评估院管理员还可在规则基础上进行必要的手动增减，并最终确认名单，兼顾规范性与灵活性。"
         "二是创新专家配置机制，系统支持一个学生关联多位专家，也支持同一专家按学校范围或指标范围承担评审任务，"
         "满足复杂场景下的专业化分工需求，通过规则化配置既避免专家任务分配失衡，也提高了评审组织的灵活性。"
         "三是创新结果查看与反馈方式，系统可实时查看专家评审进度，导出学生得分明细和指标得分情况，并基于AI智能诊断报告支持向学校反馈整体抽检结果，"
         "使结果不再停留在简单结论层面，而是能够为学校改进师范生培养、优化实践教学、强化质量保障提供数据依据和诊断建议。")

    h2("（五）推动成果可复制，解决“改革经验难推广”问题")
    body("为确保改革成果具有持久生命力和广泛影响力，项目团队在系统设计之初就注重通用性、示范性和延展性。")
    body("从平台架构看，系统采用“指标表加工具库加角色权限加流程配置加智能体”的模块化设计，三个AI智能体均可基于不同项目的指标体系和材料类型快速适配，"
         "不依赖单一学校、单一专业或单一年度任务，能够根据不同评价项目快速复制、调整和复用。")
    body("从改革路径看，该实践不是简单的信息化搬运，而是围绕评价制度执行中的堵点难点，重构标准承载方式、流程组织方式、人机协同方式和结果应用方式，"
         "形成了“数字化+智能化”双轮驱动的教育评价改革完整路径。")
    body("从应用场景看，除师范生教育教学能力认定抽检外，该模式还可推广到专业认证、教育实习评价、课程达成度监测、学生发展评价、教学质量抽样检查等工作，"
         "三个AI智能体的设计理念——预检前置、辅助评审、智能诊断——可作为通用范式迁移应用于各类教育评价场景，具有较强的迁移应用价值。")

    # ---------- 三、实施成效 ----------
    h1("三、实施成效")
    body("师范生教育教学能力认定抽检系统自上线运行以来，在效率、质量、覆盖、应用等多个维度取得了显著成效，"
         "形成了数据可量化、价值可感知、经验可推广的改革成果。")
    h2("（一）AI智能预检大幅提升审核效率与材料质量")
    body("佐证材料智能预检智能体的应用，从根本上改变了传统材料审核“逐份打开、人工判断”的低效模式。"
         "从效率维度看，单份材料平均审核时长由改革前的约20分钟下降至约10分钟，下降约50%；全市10余所高校校管累计完成500余份材料审核。"
         "从质量维度看，500余名抽检师范生中超过90%的学生材料一次性通过校审，材料一次性通过率较改革前提升约30个百分点；"
         "AI预检在校审环节即识别出约100余处多文档数据不一致问题，有效避免低质量材料流入后续评审环节。")
    h2("（二）AI评审辅助显著降低专家负担与提升意见质量")
    body("评审辅助智能体的应用，让专家从“读材料、找要点、写文字”的繁重劳动中解放出来。"
         "从效率维度看，单个学生平均评审时长由改革前的约30分钟下降至约17分钟，下降约43%。"
         "从质量维度看，AI辅助下专家撰写的评审意见的针对性、专业性显著增强。")
    h2("（三）AI智能诊断推动结果运用从“分数通报”到“质量反哺”")
    body("认定结果分析智能体的应用，彻底改变了抽检结果传统呈现方式。"
         "从效率维度看，单份学校诊断报告生成时长由改革前的约1天（人工撰写）下降至约90分钟（智能生成），效率提升明显。"
         "从学校收到诊断报告的反响看，参检学校普遍表示，新的诊断报告更便于培养方案优化、实践教学改进或薄弱指标专项强化工作，"
         "有利于形成“评价、诊断、改进、再评价”的良性闭环。从决策支撑维度看，智能诊断结果汇总形成的全市师范生教育教学能力培养质量分析报告，"
         "为重庆市教育委员会掌握全市师范生培养质量、制定优化政策提供了有力数据支撑。")

    # ---------- 四、经验总结 ----------
    h1("四、经验总结")
    body("一是坚持AI辅助、人主决策的核心原则。教育评价的专业性、敏感性、责任性，决定了AI不能替代人的专业判断。"
         "本案例三个智能体的设计始终坚守“AI识别问题、提供参考、辅助决策，人作最终判断、承担责任”的边界，"
         "既释放AI的效率价值，又守住教育评价的专业底线。这是教育领域应用AI的根本遵循。")
    body("二是坚持精准嵌入、靶向赋能的实施路径。AI赋能不能“撒胡椒面”，关键是找准业务链条上的真痛点、真堵点，以差异化AI智能体精准切入。"
         "本案例三个智能体——预检解决审核效率问题、辅助解决评审质量问题、诊断解决结果运用问题——分别针对三个不同环节的核心矛盾，"
         "避免了“为AI而AI”的泛化应用，形成了精准、务实、有效的赋能格局。")
    body("三是坚持标准先行、数据贯通的基础保障。AI智能体的能力建立在结构化指标、规范化数据和闭环化流程的基础之上。"
         "没有指标表的精细化配置、举证材料的结构化管理、评审数据的标准化沉淀，AI就成了“无米之炊”。"
         "教育评价数字化必须把基础打牢，才能让智能化真正落地。")
    body("四是坚持闭环设计、持续改进的运行机制。评价系统不仅要完成抽检任务，更要让结果看得见、问题找得到、改进有依据。"
         "本案例通过智能诊断报告将评价结果转化为改进建议，推动学校“评价、诊断、改进、再评价”的良性循环，"
         "让教育评价真正发挥质量提升的根本功能。")
    body("总体来看，师范生教育教学能力认定抽检系统以“三智能体”AI赋能体系为核心创新，破解了教育评价工作中审核效率低、评审负担重、结果运用浅等共性问题，"
         "构建了“标准入库→AI预检→规范审核→智能辅助→自动诊断→数据反馈”全流程闭环质量监测模式，"
         "形成了可复制、可推广、可持续迭代的“AI加教育评价”实践范式，对推进教育评价现代化、人工智能赋能教师教育、"
         "教育数字化转型走深走实具有较强的参考价值和示范意义。")

    out = os.path.join(BASE, "师范生教育教学能力认定抽检评价改革实践-案例文稿.docx")
    doc.save(out)

    # 统计正文中文字数（不含图题占位）
    cn = 0
    for p in doc.paragraphs:
        if "请在此处插入图片" in p.text:
            continue
        cn += len(re.findall(r'[\u4e00-\u9fff]', p.text))
    return out, cn


# ============================================================
# 文档二：附件1 征集表
# ============================================================

def set_cell(cell, text, bold=False, size=12, align=WD_ALIGN_PARAGRAPH.LEFT,
             cn_font="仿宋", valign=WD_ALIGN_VERTICAL.CENTER):
    cell.vertical_alignment = valign
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing = 1.3
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i > 0:
            p = cell.add_paragraph()
            p.alignment = align
            p.paragraph_format.line_spacing = 1.3
        run = p.add_run(line)
        set_font(run, cn_font=cn_font, size=size, bold=bold)


def shade_cell(cell, color="EAEFF7"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def build_form_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)

    pa = doc.add_paragraph()
    pa.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ra = pa.add_run("附件1")
    set_font(ra, cn_font="黑体", size=14, bold=True)

    pt = doc.add_paragraph()
    pt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pt.paragraph_format.space_after = Pt(10)
    rt = pt.add_run("教育评价领域典型实践案例征集表")
    set_font(rt, cn_font="方正小标宋简体", size=18, bold=True)

    table = doc.add_table(rows=0, cols=6)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    def add_row():
        return table.add_row().cells

    c = add_row()
    c[0].merge(c[1])
    set_cell(c[0], "案例类别", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(c[0])
    set_cell(c[2], "□ 学校发展", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(c[3], "□ 教师发展", align=WD_ALIGN_PARAGRAPH.CENTER)
    c[4].merge(c[5])
    set_cell(c[4], "☑ 学生成长", align=WD_ALIGN_PARAGRAPH.CENTER)

    c = add_row()
    c[0].merge(c[1])
    set_cell(c[0], "案例名称", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(c[0])
    merged = c[2].merge(c[3]).merge(c[4]).merge(c[5])
    set_cell(merged, CASE_NAME_ONELINE, align=WD_ALIGN_PARAGRAPH.LEFT)

    c = add_row()
    c[0].merge(c[1])
    set_cell(c[0], "案例单位", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(c[0])
    merged = c[2].merge(c[3]).merge(c[4]).merge(c[5])
    set_cell(merged, "【请填写申报单位全称】", align=WD_ALIGN_PARAGRAPH.LEFT)

    c = add_row()
    c[0].merge(c[1])
    set_cell(c[0], "类别", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(c[0])
    merged = c[2].merge(c[3]).merge(c[4]).merge(c[5])
    set_cell(merged, "□ 区（县）教科研院所　　□ 评估监测机构　　□ 高校　　□ 中小学校（含中职学校）　　□ 幼儿园\n（请按本单位性质勾选）",
             align=WD_ALIGN_PARAGRAPH.LEFT, size=11)

    c = add_row()
    set_cell(c[0], "作者姓名", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(c[0])
    set_cell(c[1], "【请填写】", align=WD_ALIGN_PARAGRAPH.CENTER)
    c[2].merge(c[3])
    set_cell(c[2], "作者单位", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(c[2])
    c[4].merge(c[5])
    set_cell(c[4], "【请填写】", align=WD_ALIGN_PARAGRAPH.CENTER)

    c = add_row()
    set_cell(c[0], "联系电话\n（手机）", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(c[0])
    set_cell(c[1], "【请填写】", align=WD_ALIGN_PARAGRAPH.CENTER)
    c[2].merge(c[3])
    set_cell(c[2], "电子邮箱", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(c[2])
    c[4].merge(c[5])
    set_cell(c[4], "【请填写】", align=WD_ALIGN_PARAGRAPH.CENTER)

    c = add_row()
    c[0].merge(c[1])
    set_cell(c[0], "案例所获\n区级及以上\n奖励情况", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(c[0])
    merged = c[2].merge(c[3]).merge(c[4]).merge(c[5])
    set_cell(merged, "无\n（如未获奖，填写“无”；如填写获奖情况，须提供相应佐证材料）",
             align=WD_ALIGN_PARAGRAPH.LEFT, size=11)

    c = add_row()
    c[0].merge(c[1])
    set_cell(c[0], "解决的\n主要问题", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(c[0])
    merged = c[2].merge(c[3]).merge(c[4]).merge(c[5])
    set_cell(merged,
             "（分条罗列关键点，150字以内）\n"
             "1.标准执行不统一，指标、评分、举证要求复杂且缺乏数字化承载，各校理解执行易偏差；\n"
             "2.审核评审效率低，材料校验仅停留在格式、数量层面，内容判断与专家评阅依赖人工，耗时长、意见易模板化；\n"
             "3.结果运用支撑弱，结果多为分数明细，缺乏结构化诊断，难以反哺培养改进。",
             align=WD_ALIGN_PARAGRAPH.LEFT, size=12)

    c = add_row()
    c[0].merge(c[1])
    set_cell(c[0], "采取的\n主要举措", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(c[0])
    merged = c[2].merge(c[3]).merge(c[4]).merge(c[5])
    set_cell(merged,
             "（分条罗列关键点，200字以内）\n"
             "研发师范生教育教学能力认定抽检系统，创新嵌入“三智能体”AI赋能体系：\n"
             "1.佐证材料智能预检，赋能学生自检与校级审核；\n"
             "2.专家评审AI辅助，自动生成结构化参考意见；\n"
             "3.认定结果智能诊断，自动生成学校诊断报告。\n"
             "同步推动标准数字化、流程闭环化、规则公平化、成果可复制，以“数字化+智能化”双轮驱动抽检质量监测。",
             align=WD_ALIGN_PARAGRAPH.LEFT, size=12)

    c = add_row()
    c[0].merge(c[1])
    set_cell(c[0], "取得的\n具体成效", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(c[0])
    merged = c[2].merge(c[3]).merge(c[4]).merge(c[5])
    set_cell(merged,
             "（分条罗列关键点，200字以内）\n"
             "1.智能预检使单份材料审核时长由约20分钟降至约10分钟（约50%），一次性通过率提升约30个百分点；\n"
             "2.AI辅助使单生评审时长由约30分钟降至约17分钟（约43%），评审意见质量提升；\n"
             "3.智能诊断报告生成由约1天降至约90分钟，结果从“分数通报”转向“质量反哺”。\n"
             "系统已支撑全市10余所高校、30余个专业、500余名师范生。",
             align=WD_ALIGN_PARAGRAPH.LEFT, size=12)

    c = add_row()
    c[0].merge(c[1])
    set_cell(c[0], "案例作者\n所属单位\n推荐意见", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(c[0])
    merged = c[2].merge(c[3]).merge(c[4]).merge(c[5])
    set_cell(merged, "\n\n\n\n单位（盖章）\n\n　　　　　　　　　　　　　　　　　年　　月　　日",
             align=WD_ALIGN_PARAGRAPH.LEFT, size=12)

    widths = [Cm(2.0), Cm(2.0), Cm(3.0), Cm(3.0), Cm(3.0), Cm(3.0)]
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths):
                cell.width = widths[idx]

    out = os.path.join(BASE, "附件1 教育评价领域典型实践案例征集表.docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    f1, cn = build_case_doc()
    f2 = build_form_doc()
    print("已生成：")
    print(" -", f1, "（正文中文字数约 %d 字）" % cn)
    print(" -", f2)
