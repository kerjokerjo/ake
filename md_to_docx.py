#!/usr/bin/env python3
"""将单元8优化稿 Markdown 转为 Word 文档。"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT


def set_run_font(run, name="宋体", size=12, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_formatted_text(paragraph, text, base_size=12, base_bold=False):
    """解析 **粗体** 与 `行内代码`。"""
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, size=base_size, bold=base_bold)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=base_size - 1)
            run.font.color.rgb = RGBColor(0x80, 0x00, 0x00)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=base_size, bold=base_bold)


def parse_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_table_separator(line):
    return bool(re.match(r"^\|?[\s\-:|]+\|?$", line.strip()))


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10.5)
                if not p.runs:
                    set_run_font(p.add_run(""), size=10.5)


def convert_md_to_docx(md_path: Path, docx_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    # 页面边距
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    i = 0
    in_code = False
    code_lines = []
    in_blockquote = False
    bq_lines = []

    def flush_blockquote():
        nonlocal in_blockquote, bq_lines
        if not bq_lines:
            return
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_after = Pt(6)
        full = "\n".join(bq_lines)
        add_formatted_text(p, full, base_size=10.5)
        in_blockquote = False
        bq_lines = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run("\n".join(code_lines))
                set_run_font(run, name="Consolas", size=10)
                run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
                in_code = False
                code_lines = []
            else:
                flush_blockquote()
                in_code = True
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # 引用块
        if stripped.startswith(">"):
            content = stripped.lstrip(">").strip()
            if not in_blockquote:
                in_blockquote = True
                bq_lines = [content]
            else:
                bq_lines.append(content)
            i += 1
            continue
        elif in_blockquote:
            flush_blockquote()

        # 空行
        if not stripped:
            i += 1
            continue

        # 分隔线
        if stripped in ("---", "***", "___"):
            i += 1
            continue

        # 表格
        if stripped.startswith("|") and "|" in stripped[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_table_separator(lines[i]):
                    table_lines.append(lines[i])
                i += 1
            if table_lines:
                rows = [parse_table_row(r) for r in table_lines]
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                style_table(table)
                for ri, row_data in enumerate(rows):
                    for ci in range(cols):
                        cell_text = row_data[ci] if ci < len(row_data) else ""
                        cell = table.rows[ri].cells[ci]
                        cell.text = ""
                        p = cell.paragraphs[0]
                        add_formatted_text(p, cell_text, base_size=10.5, base_bold=(ri == 0))
                doc.add_paragraph()
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            title = m.group(2)
            sizes = {1: 18, 2: 16, 3: 14, 4: 12}
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12 if level <= 2 else 8)
            p.paragraph_format.space_after = Pt(6)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_formatted_text(p, title, base_size=sizes.get(level, 12), base_bold=True)
            i += 1
            continue

        # 无序列表
        if re.match(r"^[-*]\s+", stripped):
            text = re.sub(r"^[-*]\s+", "", stripped)
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_text(p, text, base_size=12)
            i += 1
            continue

        # 有序列表
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_formatted_text(p, m.group(2), base_size=12)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Cm(0.74)
        add_formatted_text(p, stripped, base_size=12)
        i += 1

    flush_blockquote()
    doc.save(docx_path)
    print(f"已生成: {docx_path} ({docx_path.stat().st_size} bytes)")


if __name__ == "__main__":
    src = Path("/workspace/单元8_异常与异常处理_AI改版优化稿.md")
    dst = Path("/workspace/单元8_异常与异常处理_AI改版优化稿.docx")
    convert_md_to_docx(src, dst)
